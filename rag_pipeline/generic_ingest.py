"""Generic runtime file ingestion into ChromaDB and Neo4j.

Supported: PDF, DOCX, XLSX/XLSM, CSV, TXT, MD, JSON, PNG, JPG/JPEG, SVG.
Every readable file is chunked into ChromaDB. QODE questionnaires additionally
use the native structured graph builder. Other files use LLM entity/relation
extraction and are MERGEd into Neo4j.
"""
from __future__ import annotations
import csv, hashlib, json, logging, re
from pathlib import Path
from typing import Any

from rag_pipeline.ingest import ingest_documents, _chunk_text
from rag_pipeline.graph_builder import QODEKnowledgeGraph
from rag_pipeline.llm_client import chat

logger = logging.getLogger(__name__)
SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".md", ".json",
    ".png", ".jpg", ".jpeg", ".svg",
}

def _stable_id(prefix: str, value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")[:70] or "item"
    digest = hashlib.sha1(value.encode("utf-8", errors="ignore")).hexdigest()[:10]
    return f"{prefix}_{slug}_{digest}"

def _read_pdf(path: Path) -> str:
    import fitz
    with fitz.open(path) as doc:
        return "\n\n".join(page.get_text("text") for page in doc)

def _read_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)

def _read_excel(path: Path) -> str:
    import pandas as pd
    book = pd.ExcelFile(path)
    parts = []
    for sheet in book.sheet_names:
        df = pd.read_excel(path, sheet_name=sheet, header=None)
        parts.append(f"[Sheet: {sheet}]\n{df.fillna('').astype(str).to_csv(index=False, header=False)}")
    return "\n\n".join(parts)

def _read_csv(path: Path) -> str:
    import pandas as pd
    return pd.read_csv(path, dtype=str, keep_default_na=False).to_csv(index=False)

def _read_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    return json.dumps(data, indent=2, ensure_ascii=False)

def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")

def _read_image(path: Path) -> str:
    if path.suffix.lower() == ".svg":
        return _read_text(path)
    from PIL import Image
    import pytesseract
    with Image.open(path) as image:
        text = pytesseract.image_to_string(image).strip()
    if not text:
        raise RuntimeError(f"OCR extracted no text from {path.name}")
    return text

def read_file(path: Path) -> str:
    ext = path.suffix.lower()
    readers = {
        ".pdf": _read_pdf, ".docx": _read_docx,
        ".xlsx": _read_excel, ".xlsm": _read_excel,
        ".csv": _read_csv, ".json": _read_json,
        ".txt": _read_text, ".md": _read_text,
        ".png": _read_image, ".jpg": _read_image,
        ".jpeg": _read_image, ".svg": _read_image,
    }
    if ext not in readers:
        raise ValueError(f"Unsupported file type: {ext}")
    text = readers[ext](path).strip()
    if not text:
        raise RuntimeError(f"No text/content extracted from {path.name}")
    return text

def _documents(path: Path, text: str) -> list[dict[str, Any]]:
    docs = []
    file_key = hashlib.sha1(str(path).encode()).hexdigest()[:12]
    for idx, chunk in enumerate(_chunk_text(text)):
        docs.append({
            "id": f"runtime_{file_key}_{idx}",
            "text": chunk,
            "metadata": {
                "source": "runtime_input",
                "filename": path.name,
                "file_type": path.suffix.lower(),
                "chunk_index": str(idx),
                "diagram_type": "general",
            },
        })
    return docs

def _extract_graph(chunk: str, filename: str) -> dict:
    prompt = f"""Extract a compact knowledge graph from the supplied document chunk.
Return ONLY valid JSON with this schema:
{{"entities":[{{"name":"...", "type":"tool|role|activity|system|process|concept|metric|document"}}],
"relationships":[{{"source":"exact entity name","target":"exact entity name","type":"USES|OWNS|PERFORMS|DEPENDS_ON|PRODUCES|CONSUMES|RELATES_TO|MEASURES|MENTIONS"}}]}}
Rules: maximum 25 entities and 40 relationships; use only facts explicitly present; no markdown.
Filename: {filename}
CONTENT:
{chunk[:12000]}"""
    raw = chat([
        {"role": "system", "content": "You are a strict information extraction engine."},
        {"role": "user", "content": prompt},
    ]).strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.I | re.S)
    return json.loads(raw)

def _merge_generic_graph(graph: QODEKnowledgeGraph, path: Path, text: str) -> int:
    doc_id = _stable_id("document", path.name)
    graph._merge_node(doc_id, "document", label=path.name, filename=path.name)
    count = 0
    for chunk in _chunk_text(text):
        try:
            payload = _extract_graph(chunk, path.name)
        except Exception as exc:
            logger.warning("KG extraction skipped for a chunk of %s: %s", path.name, exc)
            continue
        entities = {}
        for item in payload.get("entities", []):
            name = str(item.get("name", "")).strip()
            if not name:
                continue
            etype = str(item.get("type", "concept")).lower()
            node_id = _stable_id(etype, name)
            entities[name.casefold()] = node_id
            graph._merge_node(node_id, etype, label=name, source_file=path.name)
            graph._merge_edge(doc_id, node_id, "MENTIONS")
            count += 1
        allowed = {"USES","OWNS","PERFORMS","DEPENDS_ON","PRODUCES","CONSUMES","RELATES_TO","MEASURES","MENTIONS"}
        for rel in payload.get("relationships", []):
            source = entities.get(str(rel.get("source", "")).strip().casefold())
            target = entities.get(str(rel.get("target", "")).strip().casefold())
            rel_type = re.sub(r"[^A-Z0-9_]", "_", str(rel.get("type", "RELATES_TO")).upper())
            if source and target and rel_type in allowed:
                graph._merge_edge(source, target, rel_type)
    return count

def is_qode_questionnaire(path: Path) -> bool:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return False
    try:
        import pandas as pd
        return "Q_Stories" in pd.ExcelFile(path).sheet_names
    except Exception:
        return False

def ingest_directory(input_dir: Path, chroma_path: str, graph_path: str) -> dict:
    files = sorted(p for p in input_dir.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS)
    if not files:
        raise RuntimeError(f"No supported input files found in {input_dir}")

    graph = QODEKnowledgeGraph()
    graph.clear_database()
    graph.build_from_pillars()
    total_docs = 0
    total_entities = 0
    processed, failed = [], []

    for path in files:
        try:
            logger.info("Processing runtime input: %s", path)
            text = read_file(path)
            docs = _documents(path, text)
            ingest_documents(docs, chroma_path)
            total_docs += len(docs)

            if is_qode_questionnaire(path):
                graph.add_excel_activities(path)
                kg_mode = "qode_structured"
            else:
                total_entities += _merge_generic_graph(graph, path, text)
                kg_mode = "llm_extracted"

            processed.append({"file": str(path), "chunks": len(docs), "kg_mode": kg_mode})
        except Exception as exc:
            logger.exception("Failed processing %s", path)
            failed.append({"file": str(path), "error": str(exc)})

    graph.save(graph_path)
    if not processed:
        raise RuntimeError(f"All input files failed: {failed}")
    return {
        "processed": processed, "failed": failed,
        "documents_upserted": total_docs,
        "generic_entities_merged": total_entities,
        "neo4j_nodes": graph.node_count,
        "neo4j_edges": graph.edge_count,
    }
